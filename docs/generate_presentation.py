import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = pptx.Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme Colors
    BG_COLOR = RGBColor(15, 17, 23)        # #0F1117 Deep Obsidian
    CARD_BG = RGBColor(26, 31, 46)         # #1A1F2E Dark Slate Surface
    CARD_BORDER = RGBColor(49, 56, 78)     # #31384E Subtle Border
    VIOLET = RGBColor(124, 58, 237)        # #7C3AED Electric Violet
    VIOLET_LIGHT = RGBColor(167, 139, 250) # #A78BFA Soft Violet
    TEXT_MAIN = RGBColor(248, 250, 252)    # #F8FAFC White/Off-white
    TEXT_MUTED = RGBColor(148, 163, 184)   # #94A3B8 Secondary Text
    TEXT_DIM = RGBColor(100, 116, 139)     # #64748B Dim Text
    EMERALD = RGBColor(16, 185, 129)       # #10B981 Success
    AMBER = RGBColor(245, 158, 11)         # #F59E0B Warning/Highlight

    def add_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()
        return bg

    def add_header(slide, badge_text, title_text, subtitle_text):
        # Badge Pill
        badge_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(4.0), Inches(0.35))
        tf_b = badge_box.text_frame
        tf_b.word_wrap = True
        p_b = tf_b.paragraphs[0]
        p_b.text = badge_text.upper()
        p_b.font.size = Pt(10)
        p_b.font.bold = True
        p_b.font.color.rgb = VIOLET_LIGHT

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.7), Inches(0.55))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(22)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_MAIN

        # Subtitle
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.7), Inches(0.45))
        tf_s = sub_box.text_frame
        tf_s.word_wrap = True
        p_s = tf_s.paragraphs[0]
        p_s.text = subtitle_text
        p_s.font.size = Pt(11)
        p_s.font.color.rgb = TEXT_MUTED

    # ════════════════════════════════════════════════════════════════════════════
    # SLIDE 1 — PROBLEM
    # ════════════════════════════════════════════════════════════════════════════
    s1 = prs.slides.add_slide(blank_layout)
    add_background(s1)
    add_header(
        s1,
        badge_text="Slide 1 — The Problem",
        title_text="Why Brands Struggle to Find & Evaluate AI-Native Creative Talent",
        subtitle_text="Traditional freelance directories are built for legacy production and fail to address generative AI workflows."
    )

    col_width = Inches(3.68)
    col_gap = Inches(0.33)
    start_x = Inches(0.8)
    cards_y = Inches(1.85)
    cards_h = Inches(4.3)

    col1_data = [
        ("Generic Marketplaces Fail AI Workflows", [
            ("Legacy Categorization", "Existing platforms categorize by 'video editor' or 'illustrator', ignoring diffusion models and neural pipelines."),
            ("No AI Stack Metadata", "Cannot filter by specific toolchains: ComfyUI, Runway Gen-3, Midjourney v6, Flux.1, or Kling AI."),
            ("Mismatch in Production Flow", "Zero support for prompt reproducibility, seed management, or iterative generative take review.")
        ]),
        ("Need for Verified Creative Capabilities", [
            ("Claims vs. Execution", "Anyone can generate lucky one-off prompts; commercial clients need consistent style and character continuity."),
            ("No Technical Proof", "Generic portfolios lack prompt breakdowns, LoRA weights, control passes, or motion vector documentation."),
            ("High Procurement Risk", "Brands risk costly project failures due to creators unable to handle targeted revisions or tight turnaround.")
        ]),
        ("Tool, Format & Commercial Requirements", [
            ("Multi-Ratio Deliverables", "Modern campaigns demand simultaneous 9:16 vertical reels, 16:9 cinematic masters, and 1:1 squares."),
            ("Complex Rights & Disclosure", "Brands require clear commercial rights transfer, explicit AI disclosure labeling, and model-training opt-outs."),
            ("Unclear Cost Structures", "Generative GPU compute and rendering compute costs are hidden or unpredictable on traditional marketplaces.")
        ])
    ]

    for i, (col_title, items) in enumerate(col1_data):
        cx = start_x + i * (col_width + col_gap)
        # Background card
        card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, cards_y, col_width, cards_h)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1)

        # Card Title
        tb = s1.shapes.add_textbox(cx + Inches(0.2), cards_y + Inches(0.2), col_width - Inches(0.4), Inches(0.6))
        p = tb.text_frame.paragraphs[0]
        p.text = col_title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = VIOLET_LIGHT

        # Card Items
        tb_items = s1.shapes.add_textbox(cx + Inches(0.2), cards_y + Inches(0.8), col_width - Inches(0.4), cards_h - Inches(0.9))
        tf = tb_items.text_frame
        tf.word_wrap = True
        for j, (heading, desc) in enumerate(items):
            p_h = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p_h.text = f"•  {heading}"
            p_h.font.size = Pt(11)
            p_h.font.bold = True
            p_h.font.color.rgb = TEXT_MAIN
            p_h.space_before = Pt(8) if j > 0 else Pt(0)

            p_d = tf.add_paragraph()
            p_d.text = desc
            p_d.font.size = Pt(10)
            p_d.font.color.rgb = TEXT_MUTED
            p_d.space_before = Pt(2)

    # Bottom summary bar
    summary_card = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.35), Inches(11.73), Inches(0.65))
    summary_card.fill.solid()
    summary_card.fill.fore_color.rgb = RGBColor(20, 24, 38)
    summary_card.line.color.rgb = VIOLET
    summary_card.line.width = Pt(1)

    stb = s1.shapes.add_textbox(Inches(0.9), Inches(6.4), Inches(11.5), Inches(0.55))
    sp = stb.text_frame.paragraphs[0]
    sp.text = "CORE PROBLEM: Brands cannot risk marketing budgets on unvetted generative claims without verified tool capability, multi-format delivery guarantees, and explicit commercial rights."
    sp.font.size = Pt(10.5)
    sp.font.bold = True
    sp.font.color.rgb = TEXT_MAIN

    # ════════════════════════════════════════════════════════════════════════════
    # SLIDE 2 — SOLUTION
    # ════════════════════════════════════════════════════════════════════════════
    s2 = prs.slides.add_slide(blank_layout)
    add_background(s2)
    add_header(
        s2,
        badge_text="Slide 2 — The Solution",
        title_text="Retake: An AI-Native Content Creator Marketplace",
        subtitle_text="A purpose-built platform connecting brands & agencies with vetted AI filmmakers, animators, and generative artists."
    )

    quad_w = Inches(5.7)
    quad_h = Inches(2.1)
    quad_x1 = Inches(0.8)
    quad_x2 = Inches(6.83)
    quad_y1 = Inches(1.85)
    quad_y2 = Inches(4.15)

    solution_pillars = [
        (quad_x1, quad_y1, "1. Targeted Two-Sided Platform", [
            ("Brands & Creative Agencies", "Search verified creators, post campaign briefs, review versioned takes, and hold milestones in INR escrow."),
            ("AI Filmmakers & Generative Artists", "Demonstrate workflow evidence, link custom LoRA pipelines, quote rates, and deliver multi-format master takes.")
        ]),
        (quad_x2, quad_y1, "2. Structured Creative Briefs", [
            ("Natural Language AI Idea Drafting", "Translates raw campaign concepts into structured objectives, formats, tools, and turnaround timelines."),
            ("Multi-Ratio & Deliverable Scoping", "Native support for simultaneous 9:16 vertical social reels, 16:9 cinematic masters, and 1:1 promo cuts.")
        ]),
        (quad_x1, quad_y2, "3. Capability-Verified Portfolios", [
            ("Verified AI Toolchain Badges", "Creators showcase verified expertise in ComfyUI, Runway Gen-3, Midjourney v6, Kling AI, and custom LoRAs."),
            ("Technical Prompt Dossiers", "Every portfolio piece includes model checkpoints, seed guidance, aspect ratios, and workflow proof.")
        ]),
        (quad_x2, quad_y2, "4. Commercial Rights & Escrow Protection", [
            ("Granular Commercial Matrix", "Pre-set usage rights (paid social, broadcast, web), territory (India/Global), term length, and model-training opt-outs."),
            ("Milestone Workflow Escrow", "Funds held securely in platform escrow and released upon milestone take approval and rights certificate issuance.")
        ]),
    ]

    for qx, qy, q_title, q_bullets in solution_pillars:
        q_card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, qx, qy, quad_w, quad_h)
        q_card.fill.solid()
        q_card.fill.fore_color.rgb = CARD_BG
        q_card.line.color.rgb = CARD_BORDER
        q_card.line.width = Pt(1)

        tb = s2.shapes.add_textbox(qx + Inches(0.2), qy + Inches(0.15), quad_w - Inches(0.4), Inches(0.4))
        p = tb.text_frame.paragraphs[0]
        p.text = q_title
        p.font.size = Pt(12.5)
        p.font.bold = True
        p.font.color.rgb = VIOLET_LIGHT

        tb_b = s2.shapes.add_textbox(qx + Inches(0.2), qy + Inches(0.55), quad_w - Inches(0.4), quad_h - Inches(0.65))
        tf = tb_b.text_frame
        tf.word_wrap = True
        for k, (title_k, text_k) in enumerate(q_bullets):
            p1 = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p1.text = f"•  {title_k}: "
            p1.font.size = Pt(10)
            p1.font.bold = True
            p1.font.color.rgb = TEXT_MAIN
            p1.space_before = Pt(4) if k > 0 else Pt(0)

            # description inline or bullet
            p1_sub = tf.add_paragraph()
            p1_sub.text = text_k
            p1_sub.font.size = Pt(9.5)
            p1_sub.font.color.rgb = TEXT_MUTED
            p1_sub.space_before = Pt(1)

    # Bottom Stat / Value Strip
    stat_box = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.45), Inches(11.73), Inches(0.55))
    stat_box.fill.solid()
    stat_box.fill.fore_color.rgb = RGBColor(22, 27, 42)
    stat_box.line.color.rgb = EMERALD
    stat_box.line.width = Pt(1)

    s_tb = s2.shapes.add_textbox(Inches(0.9), Inches(6.48), Inches(11.5), Inches(0.45))
    sp = s_tb.text_frame.paragraphs[0]
    sp.text = "CORE VALUE: Discover creators based on verified AI capability  |  Publish structured campaign briefs  |  Safeguard commercial rights"
    sp.font.size = Pt(10)
    sp.font.bold = True
    sp.font.color.rgb = TEXT_MAIN

    # ════════════════════════════════════════════════════════════════════════════
    # SLIDE 3 — PRODUCT DEMONSTRATION (Real Implementation Screenshots)
    # ════════════════════════════════════════════════════════════════════════════
    s3 = prs.slides.add_slide(blank_layout)
    add_background(s3)
    add_header(
        s3,
        badge_text="Slide 3 — Live Product Demonstration",
        title_text="Genuine Screenshots of Implemented Application Features",
        subtitle_text="Captured directly from running Next.js frontend, Express API server, and live MariaDB/MySQL database."
    )

    scr_dir = "/home/ghost/aiccm/docs/screenshots"
    shots = [
        ("1. Creator Search & Tool Filtering (/discover)", os.path.join(scr_dir, "creator_search.png"), Inches(0.8), Inches(1.85)),
        ("2. Creator Profile Dossier (/creators/arjun-patel)", os.path.join(scr_dir, "creator_profile.png"), Inches(6.8), Inches(1.85)),
        ("3. AI Portfolio Piece & Technical Metadata", os.path.join(scr_dir, "portfolio_item.png"), Inches(0.8), Inches(4.3)),
        ("4. Structured Creative Brief Builder (/briefs/new)", os.path.join(scr_dir, "brief_form.png"), Inches(6.8), Inches(4.3)),
    ]

    img_w = Inches(5.7)
    img_h = Inches(2.05)

    for label, img_path, ix, iy in shots:
        # Title above screenshot
        label_box = s3.shapes.add_textbox(ix, iy - Inches(0.24), img_w, Inches(0.25))
        lp = label_box.text_frame.paragraphs[0]
        lp.text = label
        lp.font.size = Pt(9.5)
        lp.font.bold = True
        lp.font.color.rgb = VIOLET_LIGHT

        # Border box around screenshot
        frame = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, ix - Inches(0.02), iy - Inches(0.02), img_w + Inches(0.04), img_h + Inches(0.04))
        frame.fill.solid()
        frame.fill.fore_color.rgb = CARD_BG
        frame.line.color.rgb = CARD_BORDER
        frame.line.width = Pt(1)

        if os.path.exists(img_path):
            s3.shapes.add_picture(img_path, ix, iy, width=img_w, height=img_h)

    # Caption bar
    c_box = s3.shapes.add_textbox(Inches(0.8), Inches(6.55), Inches(11.73), Inches(0.45))
    cp = c_box.text_frame.paragraphs[0]
    cp.text = "All screenshots represent functional production features: live faceted search, creator dossiers, take scrubbing, and authenticated MySQL brief persistence."
    cp.font.size = Pt(9.5)
    cp.font.italic = True
    cp.font.color.rgb = TEXT_MUTED

    # ════════════════════════════════════════════════════════════════════════════
    # SLIDE 4 — DATA MODEL AND MATCHING
    # ════════════════════════════════════════════════════════════════════════════
    s4 = prs.slides.add_slide(blank_layout)
    add_background(s4)
    add_header(
        s4,
        badge_text="Slide 4 — Data Model & Matching Architecture",
        title_text="Multi-Dimensional Metadata Representation & Match Engine",
        subtitle_text="How structured brief requirements intersect with verified creator profiles to identify the ideal talent."
    )

    dm_col_w = Inches(3.68)
    dm_col_h = Inches(4.35)
    dm_y = Inches(1.85)

    dm_cols = [
        (Inches(0.8), "CREATIVE BRIEF SCHEMA", [
            ("Deliverable Formats", "Aspect ratios (9:16 vertical reels, 16:9 cinematic masters, 1:1 square feeds) and exact durations in seconds."),
            ("AI Tool Preferences", "Keywords and declared tools (Runway Gen-3, ComfyUI, Midjourney, Kling AI, Topaz Upscale)."),
            ("Commercial License Matrix", "Permitted uses (Paid Social, Broadcast, Web), territory (India, Global), term length (12-24 mo), and AI disclosure requirement."),
            ("Budget & Turnaround Bounds", "Explicit min/max INR budgets and milestone completion deadlines.")
        ], VIOLET_LIGHT),
        (Inches(4.82), "INTELLIGENT MATCHING ENGINE", [
            ("Tool Verification Weight (35%)", "Matches brief requirements against creator badges (reviewed, evidence attached, custom LoRA proof)."),
            ("Format & Style Fit (25%)", "Evaluates creator past portfolio pieces against requested aspect ratios (9:16 vs 16:9) and visual style tags."),
            ("Budget & Pricing Fit (20%)", "Calculates whether creator priceFromInr falls within brief budget min/max range."),
            ("Reputation & Completion (20%)", "Considers verified client ratings, project completion count, and average turnaround days."),
            ("Smart Constraint Relaxation", "Dynamically suggests adjacent tools or turns off strict filters if search constraints are too narrow.")
        ], AMBER),
        (Inches(8.84), "CREATOR PROFILE SCHEMA", [
            ("Verified AI Toolchains", "Explicit tool catalog with verified status, motion chaining methods, and specialized workflows."),
            ("Portfolio Dossiers", "Versioned takes with exact prompt parameters, seed numbers, camera trajectory, and LoRA checkpoints."),
            ("Pricing & Rate Tiers", "Starting price in INR, standard turnaround days, and commercial usage surcharge policies."),
            ("Real-Time Availability", "Status indicators (Available, In Production, Booked) synced with MySQL production trackers.")
        ], EMERALD),
    ]

    for cx, c_title, c_items, c_color in dm_cols:
        card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, dm_y, dm_col_w, dm_col_h)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1)

        tb_t = s4.shapes.add_textbox(cx + Inches(0.2), dm_y + Inches(0.18), dm_col_w - Inches(0.4), Inches(0.45))
        p = tb_t.text_frame.paragraphs[0]
        p.text = c_title
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = c_color

        tb_b = s4.shapes.add_textbox(cx + Inches(0.2), dm_y + Inches(0.65), dm_col_w - Inches(0.4), dm_col_h - Inches(0.75))
        tf = tb_b.text_frame
        tf.word_wrap = True
        for k, (item_head, item_desc) in enumerate(c_items):
            p_h = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p_h.text = f"▸ {item_head}"
            p_h.font.size = Pt(10)
            p_h.font.bold = True
            p_h.font.color.rgb = TEXT_MAIN
            p_h.space_before = Pt(6) if k > 0 else Pt(0)

            p_d = tf.add_paragraph()
            p_d.text = item_desc
            p_d.font.size = Pt(9.5)
            p_d.font.color.rgb = TEXT_MUTED
            p_d.space_before = Pt(1)

    # Bottom sync text
    bot_box = s4.shapes.add_textbox(Inches(0.8), Inches(6.35), Inches(11.73), Inches(0.55))
    bp = bot_box.text_frame.paragraphs[0]
    bp.text = "Database Integration: Implemented across relational MySQL tables: users, creators, briefs, applications, and trackers with JSON document indexing."
    bp.font.size = Pt(10)
    bp.font.bold = True
    bp.font.color.rgb = TEXT_MUTED

    # ════════════════════════════════════════════════════════════════════════════
    # SLIDE 5 — RESULTS AND NEXT STEPS
    # ════════════════════════════════════════════════════════════════════════════
    s5 = prs.slides.add_slide(blank_layout)
    add_background(s5)
    add_header(
        s5,
        badge_text="Slide 5 — Results & Future Roadmap",
        title_text="Implementation Summary & Strategic Next Steps",
        subtitle_text="Delivered end-to-end functionality, marketplace differentiation, and clearly defined future enhancements."
    )

    r_w = Inches(5.7)
    r_h = Inches(4.35)
    ry = Inches(1.85)

    # Left: What is Implemented
    card_l = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), ry, r_w, r_h)
    card_l.fill.solid()
    card_l.fill.fore_color.rgb = CARD_BG
    card_l.line.color.rgb = CARD_BORDER
    card_l.line.width = Pt(1)

    tb_lt = s5.shapes.add_textbox(Inches(1.0), ry + Inches(0.2), r_w - Inches(0.4), Inches(0.4))
    p = tb_lt.text_frame.paragraphs[0]
    p.text = "DELIVERED END-TO-END WORKFLOWS"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = EMERALD

    implemented_items = [
        ("Brand Discovery & Faceted Filtering", "Live search by verified AI tool, aspect ratio, budget slider, and specialization."),
        ("Authenticated Brief Publishing", "Gated brief creation with AI Idea drafting, Rights Matrix, and direct MySQL insertion."),
        ("Capability Dossiers & Take Viewer", "Full creator profiles with verified badges, prompt parameters, and interactive take version scrubber."),
        ("Production Milestone Tracker", "Phase-gated progression from Script/Prompting to Final Master Delivery with take approval."),
        ("Escrow Wallet & Rights Licensing", "INR wallet deposits, milestone escrow locking, and automated rights certificate generation."),
        ("Strict Security & Role Gates", "Protected admin console, role-gated discovery for Brands, and authenticated sessions.")
    ]

    tb_lb = s5.shapes.add_textbox(Inches(1.0), ry + Inches(0.65), r_w - Inches(0.4), r_h - Inches(0.8))
    tf_l = tb_lb.text_frame
    tf_l.word_wrap = True
    for k, (ih, idesc) in enumerate(implemented_items):
        p_h = tf_l.paragraphs[0] if k == 0 else tf_l.add_paragraph()
        p_h.text = f"✓ {ih}: "
        p_h.font.size = Pt(10)
        p_h.font.bold = True
        p_h.font.color.rgb = TEXT_MAIN
        p_h.space_before = Pt(4) if k > 0 else Pt(0)

        p_d = tf_l.add_paragraph()
        p_d.text = idesc
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = TEXT_MUTED
        p_d.space_before = Pt(1)

    # Right: Future Improvements (Clearly Labelled)
    card_r = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.83), ry, r_w, r_h)
    card_r.fill.solid()
    card_r.fill.fore_color.rgb = CARD_BG
    card_r.line.color.rgb = CARD_BORDER
    card_r.line.width = Pt(1)

    tb_rt = s5.shapes.add_textbox(Inches(7.03), ry + Inches(0.2), r_w - Inches(0.4), Inches(0.4))
    p = tb_rt.text_frame.paragraphs[0]
    p.text = "FUTURE ROADMAP (FUTURE WORK)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = VIOLET_LIGHT

    future_items = [
        ("Automated LoRA & Pipeline Verification [Future]", "Direct API integrations with Civitai & Hugging Face to run automated reproducibility tests on creator diffusion checkpoints."),
        ("Cryptographic C2PA Provenance [Future]", "Embed tamper-proof Content Credentials metadata into final video deliverables to verify human director attribution and model provenance."),
        ("Real-Time Collaborative Video Annotations [Future]", "Time-coded canvas markup and targeted prompt revision feedback directly on the video player scrubber."),
        ("Enterprise Multi-Seat Agency Accounts [Future]", "Hierarchical agency teams with role-based budget approval limits, legal sign-offs, and multi-brand sub-workspaces.")
    ]

    tb_rb = s5.shapes.add_textbox(Inches(7.03), ry + Inches(0.65), r_w - Inches(0.4), r_h - Inches(0.8))
    tf_r = tb_rb.text_frame
    tf_r.word_wrap = True
    for k, (fh, fdesc) in enumerate(future_items):
        p_h = tf_r.paragraphs[0] if k == 0 else tf_r.add_paragraph()
        p_h.text = f"★ {fh}: "
        p_h.font.size = Pt(10)
        p_h.font.bold = True
        p_h.font.color.rgb = VIOLET_LIGHT
        p_h.space_before = Pt(6) if k > 0 else Pt(0)

        p_d = tf_r.add_paragraph()
        p_d.text = fdesc
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = TEXT_MUTED
        p_d.space_before = Pt(1)

    # Bottom Differentiation Note
    diff_box = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.35), Inches(11.73), Inches(0.65))
    diff_box.fill.solid()
    diff_box.fill.fore_color.rgb = RGBColor(22, 27, 42)
    diff_box.line.color.rgb = VIOLET
    diff_box.line.width = Pt(1)

    dtb = s5.shapes.add_textbox(Inches(0.9), Inches(6.4), Inches(11.5), Inches(0.55))
    dp = dtb.text_frame.paragraphs[0]
    dp.text = "DIFFERENTIATION FROM GENERIC MARKETPLACES: Unlike Upwork or Fiverr, Retake is architected specifically around generative diffusion pipelines, verified tool badges, iterative take scrubbers, and explicit AI commercial licensing."
    dp.font.size = Pt(9.5)
    dp.font.bold = True
    dp.font.color.rgb = TEXT_MAIN

    # Output path
    out_dir = "/home/ghost/aiccm/docs"
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, "AI-Content-Creator-Marketplace-Demo.pptx")
    prs.save(out_path)
    print(f"Presentation saved successfully to {out_path} ({os.path.getsize(out_path)} bytes)")

if __name__ == "__main__":
    build_presentation()
